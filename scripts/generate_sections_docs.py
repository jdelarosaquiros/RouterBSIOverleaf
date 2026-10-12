#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["pypandoc-binary==1.15", "python-docx==1.2.0"]
# ///
"""Export each sections/*.tex to editable Word. Run: uv run scripts/generate_sections_docs.py."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile
import zipfile

import pypandoc
from docx import Document
from docx.shared import Inches, Pt


WRAPPERS = 'archivedtables|pairedarchivedfigures|widearchivedfigures|compactwidearchivedfigure'
MACROS = r'''\newcommand{\mathds}[1]{\mathbb{#1}}
\newcommand{\rotatebox}[2]{#2}
\newcommand{\afterpage}[1]{#1}
'''


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def uncomment(text):
    """Remove TeX comments, preserving escaped percentage signs."""
    return re.sub(r'(?<!\\)((?:\\\\)*)%[^\n]*', r'\1', text)


def expand(path, root, dependencies, stack=()):
    path = path.resolve()
    if not path.is_relative_to(root):
        raise ValueError(f'Include outside paper directory: {path}')
    if path in stack:
        raise ValueError(f'Cyclic include: {path}')
    dependencies[str(path.relative_to(root))] = sha(path)
    text = uncomment(path.read_text(encoding='utf-8'))

    def include(match):
        child = root / match[1]
        if not child.suffix:
            child = child.with_suffix('.tex')
        return expand(child, root, dependencies, (*stack, path))

    return re.sub(r'\\(?:input|include)\{([^{}]+)\}', include, text)


def prepare(text, root, labels, dependencies):
    # Figure placement overrides affect PDF layout only and can confuse
    # Pandoc's LaTeX reader when it parses the corresponding float.
    text = re.sub(r'\\let\\\w+figure\\figure\s*', '', text)
    text = re.sub(r'\\renewcommand\{\\figure\}\[1\]\[\]\{\\\w+figure\[[^]]*\]\}\s*', '', text)
    # Full-width float aliases are also PDF-only; preserve their figure bodies.
    text = re.sub(r'\\expandafter\\let\\expandafter\\\w+figure\\csname (?:end)?figure\*\\endcsname\s*', '', text)
    text = re.sub(r'\\renewenvironment\{figure\}\[1\]\[\]\{\\\w+figure\[[^]]*\]\}\{\\end\w+figure\}\s*', '', text)
    text = re.sub(r'\\begin\{(?:' + WRAPPERS + r')\}(?:\[[^\]]*\])?', '', text)
    text = re.sub(r'\\end\{(?:' + WRAPPERS + r')\}', '', text)
    text = re.sub(r'\\begin\{adjustbox\}\{[^{}]*\}|\\end\{adjustbox\}', '', text)
    # Flatten simple stacked headers before Pandoc mistakes their line breaks
    # for table row boundaries (which can silently discard an entire table).
    text = re.sub(r'\\shortstack(?:\[[lcr]\])?\{([^{}]*)\}',
                  lambda m: m[1].replace(r'\\', ' '), text)
    text = text.replace(r'\begin{abstract}', r'\section*{Abstract}').replace(r'\end{abstract}', '')
    text = re.sub(r'\\bibliography\{[^}]*\}', '', text)
    text = re.sub(r'\\addlinespace(?:\[[^\]]*\])?', '', text)
    text = re.sub(r'\\cmidrule(?:\[[^\]]*\])?(?:\([^)]*\))?\{[^{}]*\}', '', text)

    def reference(match):
        key = match[1]
        if key not in labels:
            raise ValueError(f'Unresolved reference {key!r}; rebuild main.tex to update main.aux')
        return labels[key]

    def numbered_float(match):
        kind, body = match[1], match[2]
        label = re.search(r'\\label\{([^{}]+)\}', body)
        if label:
            key = label[1]
            if key not in labels:
                raise ValueError(f'Unresolved caption label {key!r}; rebuild main.tex')
            prefix = ('Figure' if kind.startswith('figure') else 'Table') + ' ' + labels[key] + ': '
            body = re.sub(r'(\\caption(?:\[[^\]]*\])?\{)', lambda m: m[1] + prefix, body, count=1)
        return r'\begin{' + kind + '}' + body + r'\end{' + kind + '}'

    text = re.sub(r'\\begin\{(figure\*?|table\*?)\}(.*?)\\end\{\1\}', numbered_float, text, flags=re.S)
    text = re.sub(r'\\ref\{([^{}]+)\}', reference, text)

    def image(match):
        path = (root / match[1]).resolve()
        if not path.is_relative_to(root):
            raise ValueError(f'Image outside paper directory: {path}')
        # Word cannot embed PDF; use the archived PNG preview of the same figure.
        if path.suffix.lower() == '.pdf':
            dependencies[str(path.relative_to(root))] = sha(path)
            path = path.with_suffix('.png')
        if not path.is_file():
            raise FileNotFoundError(f'Missing Word-compatible figure: {path}')
        dependencies[str(path.relative_to(root))] = sha(path)
        return r'\includegraphics[width=6in]{' + path.as_posix() + '}'

    return MACROS + re.sub(r'\\includegraphics(?:\[[^\]]*\])?\{([^{}]+)\}', image, text)


def document_stats(path):
    with zipfile.ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError(f'Corrupt DOCX: {path}')
    doc = Document(str(path))
    return {'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables),
            'images': len(doc.inline_shapes), 'equations': len(doc.element.xpath('//m:oMath')),
            'headings': [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]}


def format_document(path):
    doc = Document(str(path))
    section = doc.sections[0]
    # Wide source tables are readable in a landscape Word section.
    wide = any(len(t.columns) > 6 for t in doc.tables)
    section.page_width = Inches(11 if wide else 8.5)
    section.page_height = Inches(8.5 if wide else 11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    doc.styles['Normal'].font.name = 'Cambria'
    doc.styles['Normal'].font.size = Pt(11)
    available = section.page_width - section.left_margin - section.right_margin
    for table in doc.tables:
        table.autofit = False
        count = len(table.columns)
        first = Inches(1.7) if count > 6 else round(available / count)
        widths = [first] + [round((available - first) / (count - 1))] * (count - 1) if count > 1 else [available]
        for col, width in zip(table.columns, widths):
            col.width = width
        for row in table.rows:
            # Set widths by grid spans so merged dataset headers stay aligned.
            seen = set()
            grid_index = 0
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                span = cell._tc.grid_span
                cell.width = sum(widths[grid_index:grid_index + span])
                grid_index += span
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8 if len(table.columns) > 6 else 10)
    for shape in doc.inline_shapes:
        if shape.width > available:
            shape.height = round(shape.height * available / shape.width)
            shape.width = available
    doc.save(str(path))


def generate(root, output):
    sources = sorted((root / 'sections').glob('*.tex'))
    if not sources:
        raise ValueError(f'No section .tex files found in {root / "sections"}')
    dependencies = {}
    expanded = {p: expand(p, root, dependencies) for p in sources}
    labels = {}
    aux = root / 'main.aux'
    if aux.exists():
        dependencies['main.aux'] = sha(aux)
        labels = dict(re.findall(r'\\newlabel\{([^{}]+)\}\{\{([^{}]*)\}', aux.read_text()))
    citation_keys = sorted({key.strip() for t in expanded.values()
                            for group in re.findall(r'\\cite\w*\*?(?:\[[^\]]*\])*\{([^{}]+)\}', t)
                            for key in group.split(',')})
    bib = root / 'references.bib'
    if citation_keys:
        dependencies['references.bib'] = sha(bib)
    pandoc = pypandoc.get_pandoc_path()
    report = {'pandoc_version': str(pypandoc.get_pandoc_version()),
              'policy': 'Editable text, Word math and native tables; figures use PNG assets. References use main.aux numbers. Bibliography formatted by Pandoc citeproc. Empty sources get an explicit placeholder.',
              'documents': [], 'dependencies': dependencies}
    # Stage all conversions; any missing source, image, citation or warning aborts before publication.
    with tempfile.TemporaryDirectory(prefix='sections-docs-') as tmp:
        stage = Path(tmp)
        for source, original in expanded.items():
            empty = not original.strip()
            text = original
            if empty:
                text = r'\section*{' + source.stem.replace('_', ' ').title() + '}\nNo section content: this source file contains only comments or whitespace.\n'
            if source.stem == 'citations':
                text = r'\section*{References}' + '\n' + (r'\nocite{' + ','.join(citation_keys) + '}' if citation_keys else 'No citations found in the section sources.')
            prepared = prepare(text, root, labels, dependencies)
            target = stage / (source.stem + '.docx')
            command = [pandoc, '--from=latex', '--to=docx', '--standalone', '--fail-if-warnings', '-o', str(target)]
            if citation_keys:
                command += ['--citeproc', '--bibliography', str(bib)]
            result = subprocess.run(command, input=prepared, text=True, capture_output=True, cwd=root)
            if result.returncode:
                raise RuntimeError(f'{source.name}: {result.stderr.strip()}')
            format_document(target)
            report['documents'].append({'source': str(source.relative_to(root)), 'output': target.name,
                                        'empty_source': empty, 'sha256': sha(target), **document_stats(target)})
        if len(report['documents']) != len(sources):
            raise AssertionError('Section count mismatch')
        for name, digest in dependencies.items():
            if sha(root / name) != digest:
                raise RuntimeError(f'Source changed during conversion: {name}')
        output.mkdir(parents=True, exist_ok=True)
        for item in report['documents']:
            target = output / item['output']
            temporary = output / ('.' + target.name + '.tmp')
            temporary.write_bytes((stage / target.name).read_bytes())
            temporary.replace(target)
        (output / 'manifest.json').write_text(json.dumps(report, indent=2) + '\n')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--paper-dir', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--output-dir', type=Path, help='Default: PAPER/sections_docs')
    args = parser.parse_args()
    root = args.paper_dir.resolve()
    output = args.output_dir.resolve() if args.output_dir else root / 'sections_docs'
    try:
        report = generate(root, output)
    except (OSError, ValueError, RuntimeError) as exc:
        parser.exit(1, f'Export failed: {exc}\n')
    print(json.dumps({'output_dir': str(output), 'document_count': len(report['documents']), 'documents': report['documents']}, indent=2))


if __name__ == '__main__':
    main()
