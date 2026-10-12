"""Run with: uv run --with pypandoc-binary==1.15 --with python-docx==1.2.0 python -m unittest discover -s scripts -p 'test_generate_sections_docs.py' -v"""
import importlib.util
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from docx import Document

SCRIPT = Path(__file__).with_name('generate_sections_docs.py')

class SectionExportTests(unittest.TestCase):
    def test_plain_section_exports_editable_heading_text_and_math(self):
        self.assertTrue(SCRIPT.exists(), 'Section exporter has not been implemented')
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            source = root / 'sections/method.tex'
            original = r'\section{Method}' + '\n' + r'Text with \textbf{bold} and \(x_s^2\).'
            source.write_text(original)
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(root / 'sections_docs/method.docx')
            self.assertEqual(doc.paragraphs[0].text, 'Method')
            self.assertTrue(any('Text with' in p.text for p in doc.paragraphs))
            self.assertEqual(len(doc.element.xpath('//m:oMath')), 1)
            self.assertEqual(source.read_text(), original)

    def test_includes_images_citations_and_empty_placeholder(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'figs').mkdir()
            import base64
            (root / 'figs/pixel.png').write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII='))
            (root / 'main.aux').write_text(r'\newlabel{fig:test}{{2}{1}{Test}{figure.2}{}}')
            (root / 'references.bib').write_text('@article{test, author={Jane Smith}, title={Test Paper}, year={2024}, journal={Journal}}')
            (root / 'sections/intro.tex').write_text(r'\section{Introduction} See Figure~\ref{fig:test} and \citep{test}. \begin{widearchivedfigures}\input{figs/example}\end{widearchivedfigures}')
            (root / 'figs/example.tex').write_text(r'\begin{figure}\includegraphics{figs/pixel.png}\caption{A test figure.}\end{figure}')
            (root / 'sections/appendix.tex').write_text('% no appendix exists\n')
            (root / 'sections/citations.tex').write_text(r'\bibliography{references}')
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(str(root / 'sections_docs/intro.docx'))
            text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('Smith', text)
            self.assertIn('Figure\u00a02', text)
            self.assertEqual(len(doc.inline_shapes), 1)
            self.assertIn('No section content', '\n'.join(p.text for p in Document(str(root / 'sections_docs/appendix.docx')).paragraphs))
            self.assertIn('Test Paper', '\n'.join(p.text for p in Document(str(root / 'sections_docs/citations.docx')).paragraphs))
            self.assertTrue((root / 'sections_docs/manifest.json').is_file())

    def test_booktabs_rules_not_exported_as_text_and_table_fills_page(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'sections/results.tex').write_text(r'\section{Results}\begin{tabular}{lrrrrrrr} Method & A & B & C & D & E & F & G \\ \cmidrule(lr){2-3} SkillRouter & 0.123 & 0.234 & 0.345 & 0.456 & 0.567 & 0.678 & 0.789 \\ \end{tabular}')
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(str(root / 'sections_docs/results.docx'))
            table = doc.tables[0]
            self.assertEqual(table.cell(1, 0).text, 'SkillRouter')
            self.assertGreater(table.columns[0].width / 914400, 1.5)
            self.assertAlmostEqual(sum(c.width for c in table.columns) / 914400, 9.6, places=2)

    def test_booktabs_group_labels_survive_spacing_commands(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'sections/results.tex').write_text(r'\section{Results}\begin{tabular}{lr} Method & Value \\ \multicolumn{2}{l}{\textit{BM25}} \\ A & 0.123 \\ \addlinespace[2pt]\midrule\multicolumn{2}{l}{\textit{SkillSight}} \\ B & 0.456 \\ \end{tabular}')
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(str(root / 'sections_docs/results.docx'))
            self.assertTrue(any('SkillSight' in c.text for t in doc.tables for r in t.rows for c in r.cells))

    def test_tex_comments_preserve_row_break_before_percent(self):
        spec = importlib.util.spec_from_file_location('exporter', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        self.assertEqual(module.uncomment('row \\\\% comment\n10\\% value'), 'row \\\\\n10\\% value')

    def test_failed_batch_keeps_existing_docs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'sections_docs').mkdir()
            previous = root / 'sections_docs/a.docx'
            previous.write_bytes(b'existing-output')
            (root / 'sections/a.tex').write_text(r'\section{A} Valid.')
            (root / 'sections/z.tex').write_text(r'\section{Z}\includegraphics{absent.png}')
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(previous.read_bytes(), b'existing-output')
            self.assertFalse((root / 'sections_docs/z.docx').exists())

    def test_caption_numbers_match_references(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'main.aux').write_text(r'\newlabel{tab:test}{{7}{1}{A}{table.7}{}}')
            (root / 'sections/results.tex').write_text(r'\section{Results} See Table~\ref{tab:test}. \begin{table}\caption[Short]{A \textbf{nested} caption.}\label{tab:test}\begin{tabular}{lr} A & B \\ X & 1 \\ \end{tabular}\end{table}')
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(str(root / 'sections_docs/results.docx'))
            self.assertTrue(any(p.text.startswith('Table 7: A nested caption.') for p in doc.paragraphs))

    def test_figure_placement_overrides_do_not_break_float_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'figs').mkdir()
            import base64
            (root / 'figs/pixel.png').write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII='))
            (root / 'sections/ablations.tex').write_text(
                r'\begingroup\let\ablationfigure\figure\renewcommand{\figure}[1][]{\ablationfigure[!t]}'
                r'\begin{figure}[!htbp]\includegraphics{figs/pixel.png}\caption{A figure.}\end{figure}\endgroup'
            )
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(str(root / 'sections_docs/ablations.docx'))
            self.assertEqual(len(doc.inline_shapes), 1)
            self.assertTrue(any('A figure.' in p.text for p in doc.paragraphs))

    def test_deferred_figures_preserve_images_captions_and_following_prose(self):
        import base64
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'pixel.png').write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII='))
            (root / 'main.aux').write_text(r'\newlabel{fig:a}{{3}{1}}' + r'\newlabel{fig:b}{{4}{1}}')
            source = (
                r'\section{Deferred} Before. \afterpage{'
                r'\begingroup\let\backbonepagefigure\figure'
                r'\renewcommand{\figure}[1][]{\backbonepagefigure[!ht]}'
                r'\begin{figure}\includegraphics{pixel.png}\caption{Nested \textbf{caption} with \{braces\}.}\label{fig:a}\end{figure}\endgroup'
                r'\begingroup'
                r'\expandafter\let\expandafter\resultsbottomfigure\csname figure*\endcsname'
                r'\expandafter\let\expandafter\endresultsbottomfigure\csname endfigure*\endcsname'
                r'\renewenvironment{figure}[1][]{\resultsbottomfigure[!b]}{\endresultsbottomfigure}'
                r'\begin{figure}\includegraphics{pixel.png}\caption{Second caption.}\label{fig:b}\end{figure}\endgroup'
                r'} After. See Figure~\ref{fig:b}.'
            )
            (root / 'sections/deferred.tex').write_text(source)
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(root / 'sections_docs/deferred.docx')
            self.assertEqual(len(doc.inline_shapes), 2)
            text = '\n'.join(p.text for p in doc.paragraphs)
            for expected in ['Before.', 'After.', 'Figure 3: Nested caption with {braces}.', 'Figure 4: Second caption.']:
                self.assertIn(expected, text)
            self.assertEqual((root / 'sections/deferred.tex').read_text(), source)

    def test_stretched_table_remains_editable_with_group_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'sections').mkdir()
            (root / 'sections/results.tex').write_text(
                r'\section{Results}\begin{table}\caption{Routing.}'
                r'\begin{tabular*}{\linewidth}{@{\extracolsep{\fill}}lr}'
                r'Method & \shortstack{N@10\\Score} \\ \multicolumn{2}{l}{\textit{Retrievers}} \\ BM25 & 0.410 \\ '
                r'\multicolumn{2}{l}{\textit{Routers}} \\ RouterBSI & 0.728 \\ '
                r'\end{tabular*}\end{table}'
            )
            result = subprocess.run([sys.executable, str(SCRIPT), '--paper-dir', str(root)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(str(root / 'sections_docs/results.docx'))
            self.assertEqual(len(doc.tables), 1)
            rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
            self.assertEqual(rows, [['Method', 'N@10 Score'], ['Retrievers'] * 2, ['BM25', '0.410'], ['Routers'] * 2, ['RouterBSI', '0.728']])

if __name__ == '__main__':
    unittest.main()
