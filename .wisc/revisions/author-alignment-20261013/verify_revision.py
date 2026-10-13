from pathlib import Path
import hashlib,json,re,sys,zipfile
import fitz
from docx import Document
R=Path(__file__).resolve().parent
ROOT=R.parents[2]
h=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
base=json.loads((R/'baseline.json').read_text())
changed=[p for p,v in base.items() if not (ROOT/p).exists() or h(ROOT/p)!=v]
assert all(not p.endswith('.csv') for p in changed),'CSV data changed'
main_changes=[p for p in changed if p.startswith('sections/') and '/appendix/' not in p]
external={}
for rel in main_changes:
 before=(R/'backup'/rel).read_text();after=(ROOT/rel).read_text()
 assert rel=='sections/ethics_statement.tex' and after==before.replace('\\section*{Ethics Statement}','\\section*{Ethical Considerations}'),('Unexpected main-text drift',rel)
 external[rel]={'before_sha256':base[rel],'after_sha256':h(ROOT/rel),'change':'Concurrent external heading rename only; preserved, not attributed to this task.'}
warning=re.compile(r'non[\s_-]?report|analysis[\s_-]only|noncanonical',re.I)
remaining=[]
for directory in ['figs','tables','sections']:
 for p in (ROOT/directory).rglob('*'):
  if not p.is_file() or any(x.startswith('.') for x in p.relative_to(ROOT).parts):continue
  if p.suffix in {'.json','.text','.tex','.md','.txt'} and warning.search(p.read_text()):remaining.append(str(p.relative_to(ROOT)))
assert not remaining,remaining
# Numeric table bodies are invariant; caption/reference updates are separate.
tables=0
for p in (R/'backup').rglob('*.tex'):
 rel=p.relative_to(R/'backup');current=ROOT/rel
 if not current.exists():continue
 old=p.read_text();new=current.read_text()
 pattern=r'\\begin\{tabular\*?\}.*?\\end\{tabular\*?\}'
 a=re.findall(pattern,old,re.S);b=re.findall(pattern,new,re.S)
 assert a==b,str(rel)
 tables+=len(a)
# Only the known caption-bearing preview may be rebuilt.
allowed={'tables/retrieval_and_reranking_performance/preview.pdf','tables/retrieval_and_reranking_performance/preview.png'}
assert all(p in allowed for p in changed if p.startswith(('figs/','tables/')) and p.endswith(('.pdf','.png','.csv'))), 'Unexpected numeric/preview change'
for directory in ['figs','tables']:
 for p in (ROOT/directory).rglob('*.pdf'):
  text='\n'.join(str(page.get_text()) for page in fitz.open(p))
  assert not warning.search(text),str(p)
paper=fitz.open(ROOT/'main.pdf');text='\n'.join(str(p.get_text()) for p in paper)
assert not warning.search(text),'Warning remains in built paper'
# All initialization corrections really compiled; no invented token counts.
for phrase in ['All four configurations start from the same initial predicates','All component configurations use the same initial predicates','All candidate-pool configurations use the same initial predicates']:
 cleaned=re.sub(r'(?m)^\d{3,4}\s*$','',text)
 cleaned=re.sub(r'-\s*\n\s*','',cleaned)
 assert re.sub(r'\s+','',phrase) in re.sub(r'\s+','',cleaned),phrase
files=['replay_guarantee_minimal_changes.docx','protocol_details_minimal_changes.docx']
for name in files:
 p=R/name
 with zipfile.ZipFile(p) as z:assert z.testzip() is None
 d=Document(str(p));assert len(d.paragraphs)>5
 assert 'not applied' in '\n'.join(x.text for x in d.paragraphs).lower()
report={'status':'PASS','changed_baseline_files':changed,'csv_files_unchanged':sum(p.endswith('.csv') for p in base),'numeric_table_bodies_unchanged':tables,'main_proposal_only_sources_unchanged':True,'external_changes':external,'presentation_warning_matches':remaining,'paper_pages':len(paper),'paper_sha256':h(ROOT/'main.pdf'),'docx_sha256':{n:h(R/n) for n in files},'note':'Author-confirmed initialization used for prose; no experimental re-certification or new runs.'}
(R/'final_verification.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k!='changed_baseline_files'},indent=2))
