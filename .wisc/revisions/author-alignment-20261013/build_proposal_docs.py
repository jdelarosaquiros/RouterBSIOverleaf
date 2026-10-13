from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
PACK = json.loads((HERE / 'minimal_changes_proposals.json').read_text())
BY_ID = {r['id']: r for r in PACK['proposals']}


def para(doc, text, size: float = 10, bold=False, color=None, mono=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.widow_control = True
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size + 1)
    if mono:
        r.font.name = 'Liberation Mono'
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def base(title, subtitle):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin = s.bottom_margin = Inches(.65)
    s.left_margin = s.right_margin = Inches(.72)
    normal = d.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.05
    for n in ('Heading 1', 'Heading 2'):
        d.styles[n].font.name = 'Calibri'
        d.styles[n].font.color.rgb = RGBColor.from_string('17365D')
        d.styles[n].paragraph_format.space_before = Pt(8)
        d.styles[n].paragraph_format.space_after = Pt(6)
    d.styles['Heading 1'].font.size = Pt(16)
    d.styles['Heading 2'].font.size = Pt(12)
    header = s.header.paragraphs[0]
    header.text = 'RouterBSIPaper | Author-alignment proposals'
    header.style = d.styles['Caption']
    foot = s.footer.paragraphs[0]
    foot.add_run('PROPOSALS ONLY — NOT APPLIED   |   ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    for run in foot.runs:
        run.font.size = Pt(8)
    d.add_heading(title, 0)
    para(d, subtitle, 10, True, '17365D')
    para(d, 'Exact source text is quoted verbatim (LaTeX where applicable). Page references identify the saved before.pdf snapshot, not a newly compiled manuscript. All edits below are NOT APPLIED.', 9)
    return d


def source(doc, row):
    para(doc, f"Source: {row['file']}:{row['source_line_start']}–{row['source_line_end']} | before.pdf p. {row['pdf_page']}", 8.5, color='555555')


def proposal(doc, id_, title, include_reason=True):
    r = BY_ID[id_]
    text = (ROOT / r['file']).read_text()
    assert text.count(r['old']) == 1, id_
    assert text[:text.index(r['old'])].count('\n') + 1 == r['source_line_start'], id_
    doc.add_heading(title + ' — NOT APPLIED', 2)
    source(doc, r)
    if r['operation'] == 'insert_after':
        para(doc, 'Existing anchor — retain unchanged (exact source):', 9, True)
    else:
        para(doc, 'OLD — exact source fragment:', 9, True)
    para(doc, r['old'], 9)
    para(doc, 'INSERT immediately after anchor:' if r['operation'] == 'insert_after' else 'NEW — exact replacement:', 9, True, '176144')
    para(doc, r['new'], 9.5)
    if r['new_plain'] != r['new']:
        para(doc, 'Readable proposal (same wording, references expanded):', 9, True)
        para(doc, r['new_plain'], 9.5)
    if include_reason:
        para(doc, 'Why: ' + r['reason'], 9)
    evs = r['evidence']
    refs = '; '.join(f"{e['file']}:{e['line']} (snapshot p. {e['pdf_page']})" for e in evs)
    para(doc, 'Grounding: ' + refs, 8.5, color='555555')


replay = base('Replay guarantee: minimal changes', 'Document 2 | Correct inequality; narrow the scope')
para(replay, 'The wrong-direction prose remains in Related Works; the Method inequality is already correct. The proposals below correct the prose and clarify the local guarantee without changing the algorithm.', 9)
proposal(replay, '2.1', '1. Correct the direction and its local scope')
replay.add_heading('Current inequality — CORRECT / NO CHANGE', 2)
eq = PACK['correct_current_equation']
para(replay, 'Source: sections/method.tex:60 | before.pdf p. 4, printed line 255', 8.5, color='555555')
para(replay, eq['old'], 8.7)
para(replay, 'Meaning: selected average replay error is no greater than the unchanged predicate’s error on the same replay set. This is not a guarantee for every task or future test performance.', 9)
replay.add_page_break()
replay.add_heading('Scope of the replay guarantee', 1)
proposal(replay, '2.2b', '2. Limit the problem objective')
proposal(replay, '2.3', '3. State the inequality’s scope')
replay.add_page_break()
replay.add_heading('Scope of the validation gate', 1)
proposal(replay, '2.5', '4. Name the fixed comparison baseline')
proposal(replay, '2.6', '5. Align Figure 2 with strict admission (optional)')
para(replay, 'No change is proposed to the minimum-risk equation, the validation-gate equation, or the author-confirmed shared initialization.', 9, True)
replay.save(HERE / 'replay_guarantee_minimal_changes.docx')

protocol = base('Protocol details: minimal changes', 'Document 5 | Four small, appendix-grounded edits')
para(protocol, 'The top-50 candidate pool and five-versus-ten downstream budgets already appear in the main text. The proposals clarify their interpretation; they do not claim these numbers are missing.', 9)
proposal(protocol, '5.1', '1. Correct the split description')
protocol.add_page_break()
protocol.add_heading('Candidate pool and downstream budget', 1)
proposal(protocol, '5.2', '2. Separate candidates from supplied skills')
proposal(protocol, '5.3', '3. Qualify the downstream comparison')
protocol.add_page_break()
protocol.add_heading('Uncertainty reporting', 1)
proposal(protocol, '5.4', '4. Name the statistic in Figure 5')
para(protocol, 'Shared initialization is preserved exactly as clarified by the author. The change identifies population standard deviation across three runs; it neither introduces confidence intervals nor asserts statistical significance.', 9, True)
protocol.save(HERE / 'protocol_details_minimal_changes.docx')
print(json.dumps({'created': ['replay_guarantee_minimal_changes.docx', 'protocol_details_minimal_changes.docx'], 'proposal_count': len(PACK['proposals']), 'manuscript_writes': 0}, indent=2))
